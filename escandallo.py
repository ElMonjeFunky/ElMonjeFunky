#!/usr/bin/env python3
"""Interpreta el volcado de escandallo de melón (CSV exportado del ERP) y genera
un Excel con azúcar y firmeza por variedad y finca, y la firmeza por semana.

Uso:
    python3 escandallo.py VOLCADO_GLOBAL_ESCANDALLO_2026_MELON_.csv [--salida informe.xlsx]
                          [--variedades variedades.csv]

Estructura del volcado (un bloque por producto + variedad + finca, separados por
una línea "-.-.-.-"):
    PRODUCTO;VARIEDAD;...;C.FIN.DESDE;FINCA DESDE;...      cabecera del bloque
    PROBLEMAS MAYORES / PROBLEMAS MENORES                  código;defecto;%
    VALORES DE AZUCAR:                                     código;rango;%
    VALORES DE AZUCAR V:                                   código;°Brix;% de muestras
    VALORES DE COLOR:                                      código;color;%
    VALORES DE FIRMEZA + FECHA;ALBARAN;VALOR;MARCA;MIN;MAX  una fila por medida
El azúcar solo viene como distribución (% de muestras por °Brix) del bloque
entero, sin fecha; la firmeza viene medida a medida con su fecha (AAAAMMDD
con puntos de miles, ej. 20.260.605 = 05/06/2026).
"""
import argparse
import csv
import datetime as dt
import json
import re
from collections import defaultdict
from pathlib import Path
from statistics import mean

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = Path(__file__).parent
ESCALA_MIN, ESCALA_MAX = 1, 5
TIPOS = {
    "AMARILLO": "AMARILLO", "CANTALOUPE": "CANTALOUPE", "CHARENTAIS JAUNE": "CHARENTAIS",
    "GALIA": "GALIA", "GALIA L.VIDA": "GALIA LARGA VIDA", "IVORY GAYA": "IVORY GAYA",
    "PIEL SAPO": "PIEL DE SAPO", "ORANGE CANDY": "ORANGE CANDY", "SUNUP": "SUNUP",
}


def num(texto):
    return float(texto.replace(".", "").replace(",", ".")) if "," in texto else float(texto)


def tipo_y_categoria(producto):
    p = re.sub(r"^MELON\s+", "", producto.strip())
    cat = ""
    m = re.match(r"(.*?)\s+(C2|G)$", p)
    if m:
        p, cat = m.groups()
    return TIPOS.get(p, p), cat


def muestras_estimadas(pcts):
    """Nº de muestras más pequeño compatible con los porcentajes (3,33 % -> 30)."""
    for n in range(1, 501):
        if all(abs(p * n / 100 - round(p * n / 100)) < 0.06 for p in pcts):
            return n
    return None


def leer_bloques(ruta):
    lineas = ruta.read_bytes().decode("latin-1").splitlines()
    bloques, b, seccion = [], None, None
    for fila in csv.reader(lineas, delimiter=";"):
        f = [c.strip() for c in fila] + [""] * 10
        if f[0] == "PRODUCTO" and f[1] == "VARIEDAD":
            b = {"azucar": [], "firmeza": [], "rangos": {}, "revisar": []}
            bloques.append(b)
            seccion = "cabecera"
            continue
        if b is None or not any(f) or f[0].startswith("-.-"):
            continue
        if seccion == "cabecera":
            b.update(producto=f[0], variedad=f[1], cod_finca=f[4],
                     finca=re.sub(r"\s+", " ", f[5]).strip())
            seccion = None
        elif f[0].startswith("VALORES DE AZUCAR V"):
            seccion = "azucar_v"
        elif f[0].startswith("VALORES DE AZUCAR"):
            seccion = "azucar"
        elif f[0].startswith(("VALORES DE COLOR", "PROBLEMAS")):
            seccion = "otro"
        elif f[0].startswith("VALORES DE FIRMEZA"):
            seccion = "firmeza"
        elif seccion == "firmeza" and f[0] != "FECHA":
            fecha = dt.datetime.strptime(f[0].replace(".", ""), "%Y%m%d").date()
            # Escala 1-5 (las columnas MIN/MAX del volcado a veces vienen mal: 1;50, 14;5...)
            valor = num(f[2])
            if valor > ESCALA_MAX:  # p. ej. 45 en vez de 4,5: falta la coma
                b["revisar"].append((fecha, f[1], valor, "Por encima de 5 (¿falta la coma?): excluido"))
                continue
            if valor < ESCALA_MIN:
                b["revisar"].append((fecha, f[1], valor, "Por debajo de 1: incluido"))
            b["firmeza"].append((fecha, f[1], valor))
        elif seccion == "azucar_v":
            b["azucar"].append((num(f[1]), num(f[2])))
        elif seccion == "azucar":
            b["rangos"][f[1]] = num(f[2])
    return bloques


def cargar_variedades(ruta):
    if not ruta or not ruta.exists():
        return {}
    with ruta.open(encoding="utf-8-sig") as fh:
        return {r["codigo"].strip(): r["nombre"].strip() for r in csv.DictReader(fh) if r.get("nombre")}


def estadisticas(bloques):
    """Azúcar y firmeza juntando varios bloques. El azúcar de cada bloque pesa
    según su nº de muestras estimado; la firmeza junta todas las medidas."""
    conteo = defaultdict(float)
    mayor12 = n_total = 0
    for b in bloques:
        pcts = [p for _, p in b["azucar"]]
        if not pcts:
            continue
        n = muestras_estimadas(pcts) or 100
        n_total += n
        # "MAYORES DE 12" del ERP incluye el 12; se calcula de la distribución porque
        # algunos bloques usan otros rangos (ENTRE 12 Y 14, MAYORES DE 14)
        mayor12 += sum(p for v, p in b["azucar"] if v >= 12) / 100 * n
        for v, p in b["azucar"]:
            conteo[v] += p / 100 * n
    fv = [v for b in bloques for _, _, v in b["firmeza"]]
    az = ([round(sum(v * c for v, c in conteo.items()) / sum(conteo.values()), 2),
           min(conteo), max(conteo), n_total, round(mayor12 / n_total, 4)] if n_total else [None] * 5)
    fz = [len(fv), round(mean(fv), 2), min(fv), max(fv)] if fv else [0, None, None, None]
    return az + fz


# --- Excel -----------------------------------------------------------------
FUENTE = "Arial"
CAB = PatternFill("solid", fgColor="1F4E78")
BANDA = PatternFill("solid", fgColor="EEF3F8")
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
FINO = Side(style="thin", color="BFBFBF")


def hoja(wb, titulo, cabeceras, filas, formatos=None, anchos=None, primera=False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo
    ws.append(cabeceras)
    for c in ws[1]:
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        c.fill = CAB
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for i, fila in enumerate(filas, start=2):
        ws.append(fila)
        for c in ws[i]:
            c.font = Font(name=FUENTE)
            c.border = Border(bottom=FINO)
            if i % 2:
                c.fill = BANDA
    for col, fmt in (formatos or {}).items():
        for c in ws[col][1:]:
            c.number_format = fmt
    for i, ancho in enumerate(anchos or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("volcado", type=Path)
    ap.add_argument("--salida", type=Path, default=Path("informe_escandallo_melon_2026.xlsx"))
    ap.add_argument("--general", type=Path, default=Path("Datos de azucar y firmeza por melon en general.xlsx"))
    ap.add_argument("--variedades", type=Path, default=AQUI / "variedades.csv")
    args = ap.parse_args()

    semana_de = json.loads((AQUI / "semanas_2026_ia.json").read_text(encoding="utf-8"))["fecha_a_semana"]
    nombres = cargar_variedades(args.variedades)
    bloques = leer_bloques(args.volcado)
    bloques.sort(key=lambda b: (tipo_y_categoria(b["producto"]), b["variedad"], b["finca"]))

    resumen, semanal, datos_firmeza, datos_azucar, revisar = [], [], [], [], []
    for b in bloques:
        tipo, cat = tipo_y_categoria(b["producto"])
        clave = [tipo, cat, b["variedad"], nombres.get(b["variedad"], ""), b["cod_finca"], b["finca"]]
        az = b["azucar"]
        tot = sum(p for _, p in az)
        az_stats = [round(sum(v * p for v, p in az) / tot, 2), min(v for v, _ in az),
                    max(v for v, _ in az), muestras_estimadas([p for _, p in az])] if tot else [None] * 4
        mayor12 = sum(p for v, p in az if v >= 12) / tot if tot else None
        fv = [v for _, _, v in b["firmeza"]]
        fechas = [f for f, _, _ in b["firmeza"]]
        fz = [len(fv), round(mean(fv), 2), min(fv), max(fv)] if fv else [0, None, None, None]
        semanas = sorted({semana_de.get(f.strftime("%d/%m/%Y"), "SIN_SEMANA") for f in fechas})
        resumen.append(clave + az_stats + [mayor12] + fz +
                       [min(fechas) if fechas else None, max(fechas) if fechas else None, ", ".join(semanas)])

        por_semana = defaultdict(list)
        for f, alb, v in b["firmeza"]:
            sem = semana_de.get(f.strftime("%d/%m/%Y"), "SIN_SEMANA")
            por_semana[sem].append(v)
            datos_firmeza.append(clave + [f, sem, alb, v])
        for sem, vs in sorted(por_semana.items()):
            semanal.append(clave + [sem, len(vs), round(mean(vs), 2), min(vs), max(vs)])
        for v, p in az:
            datos_azucar.append(clave + [v, p / 100])
        for f, alb, v, motivo in b["revisar"]:
            revisar.append(clave + [f, alb, v, motivo])

    codigos = sorted({b["variedad"] for b in bloques})
    wb = Workbook()

    # Agrupación tipo -> variedad -> finca (las categorías C2/G se suman a su variedad)
    arbol = defaultdict(lambda: defaultdict(list))
    for b in bloques:
        arbol[tipo_y_categoria(b["producto"])[0]][b["variedad"]].append(b)
    cab_stats = ["Azúcar media (°Brix)", "Azúcar mín", "Azúcar máx", "Muestras azúcar (estim.)",
                 "% ≥ 12 °Brix", "Firmeza nº medidas", "Firmeza media", "Firmeza mín", "Firmeza máx"]
    filas_arbol, niveles, plano = [], [], []
    for tipo in sorted(arbol):
        variedades = arbol[tipo]
        todos = [b for bs in variedades.values() for b in bs]
        fincas = len({b["cod_finca"] for b in todos})
        filas_arbol.append([tipo, f"TOTAL {tipo}", "", f"{len(variedades)} variedades", fincas] + estadisticas(todos))
        niveles.append(0)
        for var in sorted(variedades):
            bs = variedades[var]
            st = estadisticas(bs)
            nf = len({b["cod_finca"] for b in bs})
            filas_arbol.append([tipo, var, nombres.get(var, ""), "", nf] + st)
            niveles.append(1)
            plano.append([tipo, var, nombres.get(var, ""), nf] + st)
            for b in sorted(bs, key=lambda b: b["finca"]):
                cat = tipo_y_categoria(b["producto"])[1]
                finca = f"{b['finca']} ({b['cod_finca']})" + (f" [{cat}]" if cat else "")
                filas_arbol.append([tipo, var, nombres.get(var, ""), finca, 1] + estadisticas([b]))
                niveles.append(2)

    fmt_stats = {"F": "0.00", "J": "0%", "L": "0.00"}
    pv = hoja(wb, "Por variedad", ["Tipo melón", "Variedad (código)", "Nombre variedad", "Finca",
                                   "Nº fincas"] + cab_stats, filas_arbol, fmt_stats,
              [17, 22, 18, 46, 8, 11, 9, 9, 11, 9, 10, 10, 9, 9], primera=True)
    TIPO_FILL, VAR_FILL = PatternFill("solid", fgColor="BDD7EE"), PatternFill("solid", fgColor="E2EFDA")
    pv.sheet_properties.outlinePr.summaryBelow = False
    for i, nivel in enumerate(niveles, start=2):
        if nivel == 0:
            for c in pv[i]:
                c.fill, c.font = TIPO_FILL, Font(name=FUENTE, bold=True)
        elif nivel == 1:
            for c in pv[i]:
                c.fill, c.font = VAR_FILL, Font(name=FUENTE, bold=True)
            pv.row_dimensions[i].outlineLevel = 1
        else:
            for c in pv[i]:
                c.fill = PatternFill(fill_type=None)
            pv[f"D{i}"].alignment = Alignment(indent=2)
            pv.row_dimensions[i].outlineLevel = 2
    hoja(wb, "Resumen variedades", ["Tipo melón", "Variedad (código)", "Nombre variedad", "Nº fincas"]
         + cab_stats, plano, {"E": "0.00", "I": "0%", "K": "0.00"},
         [17, 18, 18, 8, 11, 9, 9, 11, 9, 10, 10, 9, 9])
    comunes = ["Tipo melón", "Cat.", "Variedad (código)", "Nombre variedad", "Cód. finca", "Finca"]
    anchos_comunes = [17, 6, 16, 18, 13, 34]

    ws = hoja(wb, "Variedad x Finca", comunes + [
        "Azúcar media (°Brix)", "Azúcar mín", "Azúcar máx", "Muestras azúcar (estim.)", "% ≥ 12 °Brix",
        "Firmeza nº medidas", "Firmeza media", "Firmeza mín", "Firmeza máx",
        "Primera fecha", "Última fecha", "Semanas con firmeza"], resumen,
        {"G": "0.00", "K": "0%", "M": "0.00", "P": "dd/mm/yyyy", "Q": "dd/mm/yyyy"},
        anchos_comunes + [11, 9, 9, 11, 9, 10, 10, 9, 9, 12, 12, 40])
    semanal_ws = hoja(wb, "Firmeza semanal", comunes + ["Semana", "Nº medidas", "Firmeza media",
                                                        "Firmeza mín", "Firmeza máx"], semanal,
                      {"I": "0.00"}, anchos_comunes + [9, 10, 10, 10, 10])
    fz_ws = hoja(wb, "Datos firmeza", comunes + ["Fecha", "Semana", "Albarán", "Firmeza"],
                 datos_firmeza, {"G": "dd/mm/yyyy"}, anchos_comunes + [12, 9, 9, 9])
    az_ws = hoja(wb, "Datos azúcar", comunes + ["°Brix", "% muestras"], datos_azucar,
                 {"H": "0.00%"}, anchos_comunes + [9, 11])
    hoja(wb, "Revisar", comunes + ["Fecha", "Albarán", "Firmeza", "Motivo"], revisar,
         {"G": "dd/mm/yyyy"}, anchos_comunes + [12, 9, 9, 42])
    var_ws = hoja(wb, "Variedades", ["codigo", "nombre"], [[c, nombres.get(c, "")] for c in codigos],
                  anchos=[18, 30])
    for c in var_ws["B"][1:]:
        c.fill = AMARILLO

    notas = wb.create_sheet("Notas")
    for linea in [
        "Cómo se ha interpretado el volcado",
        f"Origen: {args.volcado.name} — {len(bloques)} bloques (uno por producto + variedad + finca).",
        "'Por variedad': tipo de melón -> cada variedad (fila verde, todas sus fincas juntas) -> el detalle de cada finca. Los botones 1/2/3 de la izquierda pliegan/despliegan.",
        "  Al juntar fincas, el azúcar de cada finca pesa según su nº de muestras y la firmeza junta todas las medidas. C2 y G se suman a su variedad (se indican entre corchetes en la finca).",
        "Tipo melón = PRODUCTO sin 'MELON'. Cat. = sufijo del producto (C2 = categoría 2; G = sin confirmar).",
        "Variedad = código del volcado. El nombre sale de la hoja 'Variedades': se rellena desde variedades.csv (codigo,nombre); al completarlo y volver a ejecutar escandallo.py aparece en todas las hojas.",
        "AZÚCAR: el volcado solo da el % de muestras por cada valor de °Brix (sección 'VALORES DE AZUCAR V') para toda la temporada del bloque, SIN FECHA.",
        "  Media = media ponderada por ese %. Mín/Máx = valores con % > 0. Muestras (estim.) = nº mínimo de muestras compatible con los % (3,33 % -> 30).",
        "  Por eso el azúcar NO se puede dar por semana con este volcado.",
        "% ≥ 12 °Brix = % de muestras con 12 °Brix o más (equivale al rango MAYORES DE 12 del ERP).",
        "FIRMEZA: una fila por medida con su fecha (20.260.605 = 05/06/2026) y albarán. Media/mín/máx calculados sobre todas las medidas.",
        "  Escala de firmeza 1-5 (columnas MIN/MAX del volcado; en unas pocas filas vienen mal escritas y se ignoran).",
        "  Medidas por encima de 5 (23, 27, 33, 41, 45: falta la coma) se excluyen; las de menos de 1 se mantienen. Todas aparecen en la hoja 'Revisar'.",
        "Bloques sin datos de azúcar o sin medidas de firmeza aparecen con esas columnas vacías.",
        "Semanas ISO (lunes a domingo) según semanas_2026_ia.json.",
        "Los valores de las hojas de resumen están calculados por escandallo.py a partir de las hojas 'Datos firmeza' y 'Datos azúcar'.",
    ]:
        notas.append([linea])
    notas["A1"].font = Font(name=FUENTE, bold=True, size=13)
    for c in notas["A"][1:]:
        c.font = Font(name=FUENTE)
    notas.column_dimensions["A"].width = 140

    wb.save(args.salida)
    # Versión resumida: solo el desglose por variedad
    for nombre in wb.sheetnames:
        if nombre not in ("Por variedad", "Resumen variedades", "Variedades", "Notas"):
            del wb[nombre]
    wb["Por variedad"].sheet_view.zoomScale = 90
    wb.save(args.general)
    print(f"{len(bloques)} bloques, {len(datos_firmeza)} medidas de firmeza, {len(codigos)} variedades -> {args.salida}")


if __name__ == "__main__":
    main()
